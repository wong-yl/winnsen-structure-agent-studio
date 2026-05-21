import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "data"

CASE_DIR = Path(r"D:\机械结构工程师智能体\outputs\16029_case1_10door")
PARAM_CARD = CASE_DIR / "16029标准柜试改1_2-10十门柜无脑制作参数卡.xlsx"
TASK_BOOK = CASE_DIR / "16029标准柜试改1_10门2列每列5门改型任务单.xlsx"
PRACTICE_REPORT = CASE_DIR / "16029标准柜试改1_D盘练习副本_2-10十门柜自动复核报告.xlsx"
PRACTICE_SOURCE_ASSEMBLY = Path(
    r"D:\机械结构工程师智能体\work\16029_练习副本_20260429\1.工程图\标准寄存柜1917×1000×550(总装配).SLDASM"
)
MANUAL_RUN_DIR = ROOT / "workers" / "manual_runs"

RECIPE_JSON = DATA_DIR / "solidworks_16029_10door_mutator_recipe.json"
RECIPE_MD = DATA_DIR / "solidworks_16029_10door_mutator_recipe.md"


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def workbook_rows(path: Path, sheet_name: str) -> list[dict[str, str]]:
    if not path.exists():
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [cell_text(value) for value in rows[0]]
    parsed: list[dict[str, str]] = []
    for row in rows[1:]:
        item = {
            headers[index]: cell_text(value)
            for index, value in enumerate(row)
            if index < len(headers) and headers[index]
        }
        if any(item.values()):
            parsed.append(item)
    return parsed


def key_value_rows(path: Path, sheet_name: str, key_header: str, value_header: str) -> dict[str, str]:
    return {
        row[key_header]: row.get(value_header, "")
        for row in workbook_rows(path, sheet_name)
        if row.get(key_header)
    }


def file_info(path: Path | None) -> dict[str, Any]:
    exists = bool(path and path.exists())
    stat = path.stat() if path and exists else None
    return {
        "path": str(path) if path else "",
        "exists": exists,
        "sizeBytes": stat.st_size if stat else None,
        "updatedAt": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat() if stat else None,
    }


def latest_matching_file(pattern: str) -> Path | None:
    matches = [path for path in MANUAL_RUN_DIR.glob(pattern) if path.is_file() and not path.name.startswith("~$")]
    if not matches:
        return None
    return max(matches, key=lambda path: path.stat().st_mtime)


def build_payload() -> dict[str, Any]:
    parameter_card = key_value_rows(PARAM_CARD, "一页参数卡", "项目", "直接执行值")
    calculation_basis = workbook_rows(PARAM_CARD, "AI计算依据")
    engineer_steps = workbook_rows(PARAM_CARD, "结构工程师执行清单")
    review_conclusion = key_value_rows(PRACTICE_REPORT, "复核结论", "项目", "结论")
    quantity_checks = workbook_rows(PRACTICE_REPORT, "核心数量复核")
    change_details = workbook_rows(PRACTICE_REPORT, "变化明细")
    control_points = workbook_rows(TASK_BOOK, "控制点修改建议")
    ten_door_smoke_output = latest_matching_file("QA-16029-10DOOR-FAST-CLONE-*/*.SLDASM")

    return {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "modelFamily": "16029 standard locker",
        "target": "10-door same-envelope variant",
        "outerSizeMm": {"width": 1000, "height": 1917, "depth": 550},
        "sourceFiles": {
            "parameterCard": file_info(PARAM_CARD),
            "taskBook": file_info(TASK_BOOK),
            "practiceReviewReport": file_info(PRACTICE_REPORT),
            "practiceSourceAssembly": file_info(PRACTICE_SOURCE_ASSEMBLY),
        },
        "keyParameters": {
            "layout": parameter_card.get("门布局", "2列 x 每列5门 = 10门"),
            "doorType": parameter_card.get("门型定义", "2/10"),
            "baseDoorZoneHeightMm": parameter_card.get("原门区高度基准", "1830.00mm"),
            "unitHeightMm": parameter_card.get("新高度单位", "183.00mm"),
            "doorPitchMm": parameter_card.get("新门阵列节距", "366.00mm"),
            "doorPanelHeightMm": parameter_card.get("新门板标注高", "359.00mm"),
            "doorPanelWidthMm": parameter_card.get("门板宽", "437.00mm"),
        },
        "mutatorOperations": [
            {
                "operation": "derive_2_10_door_panel",
                "detail": "Copy 2/12 door-panel source to 2/10, set panel height to 359.00mm and keep width near 437.00mm.",
            },
            {
                "operation": "derive_2_10_door_weld_and_assembly",
                "detail": "Copy 2/12 door weldment and assembly to 2/10 names, replace the inner panel with the 2/10 panel, keep ZJA-S500 lock hardware.",
            },
            {
                "operation": "update_door_pattern",
                "detail": "Set cabinet door pattern count from 6 to 5 and pitch from 305.00mm to 366.00mm.",
            },
            {
                "operation": "update_frame_and_shelf_patterns",
                "detail": "Set door-frame divider and shelf patterns from 11 x 152.50mm to 9 x 183.00mm; recompute skipped instances.",
            },
        ],
        "solidworksTemplateClone": {
            "supportedDoorCountsAfterThisRecipe": [10, 12],
            "runnerScript": r"D:\winnsen_cad_workspace\scripts\sw_clone_16029_baseline_template.js",
            "tenDoorSourceAssembly": str(PRACTICE_SOURCE_ASSEMBLY),
            "status": "usable_as_template_clone_native_save_smoke_passed_packgo_pending",
            "tenDoorNativeSaveSmokeOutput": file_info(ten_door_smoke_output),
        },
        "reviewConclusion": review_conclusion,
        "quantityChecks": quantity_checks,
        "changeDetails": change_details,
        "calculationBasis": calculation_basis,
        "engineerSteps": engineer_steps,
        "controlPoints": control_points,
        "boundaries": [
            "10-door is available as a SolidWorks practice-template clone, not yet as a fully algorithmic arbitrary door-count mutator.",
            "Do not overwrite the original 12-door standard template.",
            "Native SolidWorks save smoke has passed for the 10-door template clone; Pack-and-Go is still required before independent handoff.",
        ],
    }


def write_markdown(payload: dict[str, Any]) -> None:
    parameters = payload["keyParameters"]
    source = payload["sourceFiles"]["practiceSourceAssembly"]
    lines = [
        "# 16029 10门变体配方",
        "",
        f"- Generated at: `{payload['generatedAt']}`",
        f"- Status: `{payload['solidworksTemplateClone']['status']}`",
        f"- Practice SolidWorks source: `{source['path']}`",
        f"- Source exists: `{source['exists']}`",
        f"- 10-door native save smoke: `{payload['solidworksTemplateClone']['tenDoorNativeSaveSmokeOutput']['path']}`",
        "",
        "## 关键参数",
        "",
        "| Item | Value |",
        "| --- | --- |",
    ]
    for key, value in parameters.items():
        lines.append(f"| {key} | {value} |")
    lines.extend(["", "## SolidWorks 修改动作", "", "| Operation | Detail |", "| --- | --- |"])
    for item in payload["mutatorOperations"]:
        lines.append(f"| {item['operation']} | {item['detail']} |")
    lines.extend(["", "## 数量复核", "", "| Check | Original | Current | Delta | Status | Note |", "| --- | ---: | ---: | ---: | --- | --- |"])
    for row in payload["quantityChecks"]:
        lines.append(
            "| {check} | {old} | {current} | {delta} | {status} | {note} |".format(
                check=row.get("检查项", ""),
                old=row.get("原模板数量", ""),
                current=row.get("当前数量", ""),
                delta=row.get("变化", ""),
                status=row.get("状态", ""),
                note=row.get("说明", ""),
            )
        )
    lines.extend(["", "## 当前边界", ""])
    for item in payload["boundaries"]:
        lines.append(f"- {item}")
    RECIPE_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    payload = build_payload()
    RECIPE_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    write_markdown(payload)
    print(f"wrote {RECIPE_JSON}")
    print(f"wrote {RECIPE_MD}")


if __name__ == "__main__":
    main()
