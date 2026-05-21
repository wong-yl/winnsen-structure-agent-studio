from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from build_winnsen_leadership_pdf import prepare_assets
from build_winnsen_progress_report import BLUE, DOCS, LIGHT_BLUE, LIGHT_ORANGE, MODEL_DIR, MUTED, ORANGE, TEXT, load_data


OUT_DOCX = DOCS / "Winnsen结构智能体项目领导汇报_20260519.docx"


def set_run_font(run, size: float = 10.5, bold: bool = False, color: str = TEXT) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.replace("#", ""))
    tc_pr.append(shd)


def cell_text(cell, text: str, bold: bool = False, color: str = TEXT, size: float = 9.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_title(doc: Document, title: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    set_run_font(r, size=24, bold=True, color=BLUE)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(16)
        r2 = p2.add_run(subtitle)
        set_run_font(r2, size=11.5, color=MUTED)


def add_section(doc: Document, title: str, subtitle: str | None = None) -> None:
    if len(doc.paragraphs) > 0:
        doc.add_page_break()
    add_title(doc, title, subtitle)


def add_para(doc: Document, text: str, size: float = 10.5, color: str = TEXT, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(5)
        marker = p.add_run("• ")
        set_run_font(marker, size=10.5, bold=True, color=ORANGE)
        r = p.add_run(item)
        set_run_font(r, size=10.5, color=TEXT)


def add_table(doc: Document, rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.rows[row_index].cells[col_index]
            if row_index == 0:
                shade_cell(cell, BLUE)
                cell_text(cell, value, bold=True, color="FFFFFF", size=9.5)
            else:
                if row_index % 2 == 1:
                    shade_cell(cell, "F7F9FC")
                cell_text(cell, value, size=9.2)
            cell.width = Cm(widths[col_index])
    doc.add_paragraph()


def add_image(doc: Document, path: Path, width_cm: float = 24.0) -> None:
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def add_callout(doc: Document, title: str, body: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade_cell(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_run_font(r, size=11.5, bold=True, color=BLUE)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.2, color=TEXT)
    doc.add_paragraph()


def build_docx() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    data = load_data()
    assets = prepare_assets(data)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10)

    add_title(doc, "Winnsen 硬件结构知识与钣金模型生成平台 MVP", "项目内容及进度汇报 | 面向领导汇报 | 2026-05-19")
    add_callout(
        doc,
        "当前阶段",
        "已完成 16029 / 1000W x 1917H x 550D / 10门 FreeCAD 规则参考模型，并通过基础校验。",
    )
    add_table(
        doc,
        [
            ["汇报重点", "说明"],
            ["工程价值", "沉淀门数变化规则，减少结构工程师重复建模和前期判断时间。"],
            ["当前成果", "平台链路 + 数据边界 + 双 CAD 路线 + 首个可校验模型。"],
            ["下一步", "生成 12/14 门对照样本，形成同外形门数变化规则表。"],
        ],
        [5, 18.5],
    )

    add_section(doc, "阶段总览", "3 分钟汇报结论")
    add_para(doc, "项目已经从页面原型推进到“有一个可校验规则生成模型”的阶段；后续不再堆模型库存，而是用代表样本提炼可复用参数化规则。", size=12, bold=True)
    add_table(
        doc,
        [
            ["维度", "当前结论"],
            ["业务问题", "标准模型多、尺寸变体多、重复建模耗时，结构规则难复用。"],
            ["当前成果", f"16029 / 1000W x 1917H x 550D / {data.door_count}门工程参考模型已生成并校验通过。"],
            ["技术路线", "FreeCAD 先验证规则；SolidWorks 作为工程交接和复核通道。"],
            ["项目边界", "当前不是生产图纸；正式 SolidWorks 图纸、DXF、BOM 仍需工程流程放行。"],
        ],
        [4.5, 20],
    )

    add_section(doc, "前期工作转化价值", "前期不是白干，而是把“能不能做”变成“怎么稳定做”")
    add_para(doc, "前十几天的大部分工作不是最终模型本身，而是把数据、入口、生成链路、校验和风险边界摸清。领导汇报里需要把这些工作翻译成可理解的阶段成果。")
    add_table(
        doc,
        [
            ["前期工作", "转化成的项目资产", "对后续的价值"],
            ["五个 MVP 页面", "项目总览、数据录入、规则成熟度、可生成模型、待确认项", "让项目从散乱脚本变成可管理平台"],
            ["CAD 数据资产梳理", "明确 16029 柜体族、标准模型、补充素材、输出目录", "后续生成不再凭空猜测，有来源和边界"],
            ["SolidWorks / FreeCAD 双入口", "建立两条 CAD 路线和任务记录", "SolidWorks 不顺时，FreeCAD 仍能推进规则验证"],
            ["SolidWorks 试错", "发现自动装配错位、窗口过多、资源压力等风险", "避免继续把时间耗在不稳定路径上"],
            ["校验机制", "数量校验、bbox、几何完整性、工程交付说明", "模型不只看起来像，还要有数据证明可复核"],
            ["路线收敛", "从“克隆模型”调整为“规则生成 + 工程交接”", "把项目目标拉回工程师提效，而不是堆展示页面"],
        ],
        [5.5, 9, 10],
    )
    add_callout(doc, "汇报口径", "前期工作应表述为“基础能力建设和风险收敛”，不是最终模型数量。", fill=LIGHT_ORANGE)

    add_section(doc, "项目背景与目标", "项目解决的问题与当前目标")
    add_bullets(
        doc,
        [
            "重复建模：不同门数、不同门高、不同五金组合需要反复调整。",
            "规则分散：经验藏在历史模型、文件夹和工程师个人习惯中。",
            "交接成本：没有统一的参数、校验和交付格式，工程复核效率低。",
        ],
    )
    add_table(
        doc,
        [
            ["目标", "说明", "当前状态"],
            ["规则沉淀", "把门数、门距、层板、锁位、铰链等变化关系转成可复用规则。", "进行中"],
            ["模型生成", "按参数生成工程参考模型，输出 STEP/FCStd 和校验报告。", "10门已通过"],
            ["工程提效", "让结构工程师先拿到可复核模型和参数表，再进入正式图纸。", "已形成首个样本"],
        ],
        [4.2, 15, 5],
    )

    add_section(doc, "平台能力", "MVP 已具备可操作入口和双 CAD 路线")
    add_image(doc, assets["platform"], 24.5)

    add_section(doc, "总体方案", "从标准模型到工程参考模型")
    add_image(doc, assets["flow"], 24.5)

    add_section(doc, "数据资产与范围", "当前先聚焦 16029 柜体族，不泛化到所有产品")
    add_table(
        doc,
        [
            ["项目", "说明"],
            ["主数据工作区", r"D:\机械结构工程师智能体：历史 CAD、STEP、SolidWorks 证据、组件/transform/bbox 信息。"],
            ["补充素材", r"C:\Users\Administrator\Desktop\参数化模板素材：后续补齐标准样本和特殊门型。"],
            ["当前交付包", str(MODEL_DIR)],
            ["当前优先级", "先跑通 16029 同外形 10/12/14 门变化规则，再扩展宽度和特殊门型。"],
        ],
        [5.2, 19.2],
    )

    add_section(doc, "已完成模型", "16029 / 1000W x 1917H x 550D / 10门")
    add_image(doc, assets["model"], 24.5)

    add_section(doc, "规则种子", "当前已沉淀的 16029 10门规则种子")
    add_table(
        doc,
        [
            ["规则项", "当前值", "工程意义"],
            ["外形目标", "1000W x 1917H x 550D", "同外形门数变化的第一组样本"],
            ["门数布局", f"{data.door_count}门 / 2列 x {data.rows_per_column}行", "用于推导 12/14 门对照规则"],
            ["单门尺寸", f"{data.door_width:.0f}W x {data.door_height:.0f}H mm", "门高随门数变化重算"],
            ["门距", f"{data.door_pitch:.0f} mm", "驱动门板、层板、横隔、锁位"],
            ["数量校验", "10门模块 / 10锁孔 / 8层板 / 8横隔", "基础数量关系已通过"],
            ["几何完整性", "invalid shape = 0", "FCStd 基础完整性检查通过"],
        ],
        [5.2, 7.5, 11.8],
    )
    add_callout(doc, "结论", "当前不是在提前做所有变体，而是在用代表样本提炼可复用生成规则。", fill=LIGHT_ORANGE)

    add_section(doc, "校验结果", "基础完整性和数量规则均通过")
    add_image(doc, assets["validation"], 24.5)

    add_section(doc, "风险与调整", "SolidWorks 保留为工程交接和复核通道")
    add_table(
        doc,
        [
            ["风险", "表现", "处理策略"],
            ["SolidWorks API 工程化风险", "早期自动装配出现错乱，不能作为当前主线成果。", "先用 FreeCAD 验证规则，再回到 SolidWorks 做单模型验证。"],
            ["电脑资源压力", "同时打开或生成多个模型会拖慢机器。", "后续每次只跑一个 CAD 任务，不并行跑。"],
            ["生产交付边界", "STEP/FCStd 可参考，但不是完整 SolidWorks 特征树。", "明确标注工程参考模型，正式图纸/DXF/BOM 后续放行。"],
        ],
        [5.8, 8.8, 9.8],
    )
    add_callout(doc, "路线调整后的判断", "当前最有价值的不是继续克隆已有模型，而是把 16029 门数变化规则跑通，让工程师拿到能复核、能继续建图的参数和参考模型。", fill=LIGHT_ORANGE)

    add_section(doc, "下一阶段计划", "先验证规则，再进入 SolidWorks 工程交接")
    add_image(doc, assets["roadmap"], 24.5)
    add_para(doc, "执行原则：每次只跑一个 CAD 任务；普通插话不重置阶段；电脑资源异常、模型明显错乱、用户明确叫停时才打断。")

    add_section(doc, "阶段结论", "当前可向领导汇报的结论")
    add_table(
        doc,
        [
            ["结论", "说明"],
            ["已完成一个有真实数据支撑的工程参考模型", f"16029 / {data.door_count}门 / 1000W x 1917H x 550D，基础校验 PASS。"],
            ["项目方向已经从“页面展示”转向“规则生成”", "模型、参数、校验和交付包开始形成闭环。"],
            ["下一阶段目标清晰", "先完成 10/12/14 门同外形规则表，再决定 SolidWorks 工程交接方式。"],
        ],
        [8, 16.5],
    )

    doc.core_properties.title = "Winnsen结构智能体项目领导汇报"
    doc.core_properties.subject = "MVP 项目内容及进度汇报"
    doc.core_properties.author = "Winnsen Structure Agent Studio"
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build_docx()
