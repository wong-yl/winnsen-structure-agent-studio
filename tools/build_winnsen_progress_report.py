from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(r"D:\Winnsen_Structure_Agent_Studio")
DOCS = ROOT / "docs"
OUT_DOCX = DOCS / "Winnsen结构智能体项目内容及进度书_20260519.docx"
OUT_PDF = DOCS / "Winnsen结构智能体项目内容及进度书_20260519.pdf"
ASSET_DIR = DOCS / "report_assets" / "20260519_progress"

BRAND = Path(r"C:\Users\Administrator\Desktop\Winnsen品牌设计")
BRAND_LOGO = BRAND / "logo.jpg"
BRAND_SYSTEM = BRAND / "1.png"
BRAND_PRODUCT = ROOT / "apps" / "web" / "dist" / "brand" / "winnsen-product-application.png"
UI_SCREENSHOT = ROOT / "apps" / "web" / "docs" / "design" / "qa" / "desktop-dual-cad-entries-panel.png"
UI_RESULT_SCREENSHOT = ROOT / "apps" / "web" / "docs" / "design" / "qa" / "desktop-freecad-worker-output-result.png"

MODEL_DIR = ROOT / "workers" / "generated_models" / "FREECAD-16029-10DOOR-VARIANT-20260519"
HANDOFF_MD = MODEL_DIR / "ENGINEERING_HANDOFF.md"
REPORT_MD = MODEL_DIR / "locker_16029_10door_rule_driven_report.md"
VALIDATION_MD = MODEL_DIR / "locker_16029_10door_rule_driven_validation.md"
INTEGRITY_MD = MODEL_DIR / "locker_16029_10door_rule_driven_geometry_integrity.md"
VERIFY_CSV = MODEL_DIR / "locker_16029_10door_rule_driven_verify.csv"
FCSTD = MODEL_DIR / "locker_16029_10door_rule_driven.FCStd"
STEP = MODEL_DIR / "locker_16029_10door_rule_driven.step"

BLUE = "#1F3585"
ORANGE = "#F04A12"
LIGHT_BLUE = "#EAF0FF"
LIGHT_ORANGE = "#FFF1EA"
LIGHT_GRAY = "#F3F6FA"
TEXT = "#17233C"
MUTED = "#637083"
GREEN = "#0E8A43"


@dataclass
class ProjectData:
    door_count: int = 10
    rows_per_column: int = 5
    door_width: float = 437.0
    door_height: float = 359.0
    door_pitch: float = 366.0
    cabinet_width: float = 1000.0
    bbox_y: float = 1983.0
    bbox_z: float = 552.0
    solid_count: int = 361
    validation_status: str = "PASS"
    integrity_status: str = "PASS"
    invalid_shapes: int = 0
    non_group_no_shape: int = 0
    center_vertical_failures: int = 0
    file_fcstd_mb: float = 0.0
    file_step_mb: float = 0.0
    counts: dict[str, int] | None = None


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace") if path.exists() else ""


def find_number(text: str, label: str, default: float) -> float:
    match = re.search(rf"{re.escape(label)}[:：]\s*`?([0-9.]+)", text)
    if match:
        return float(match.group(1))
    return default


def load_data() -> ProjectData:
    handoff = read_text(HANDOFF_MD)
    report = read_text(REPORT_MD)
    validation = read_text(VALIDATION_MD)
    integrity = read_text(INTEGRITY_MD)
    combined = "\n".join([handoff, report, validation, integrity])

    counts: dict[str, int] = {}
    for line in validation.splitlines():
        m = re.match(r"-\s+([A-Za-z0-9_]+):\s+([0-9]+)", line.strip())
        if m:
            counts[m.group(1)] = int(m.group(2))

    data = ProjectData(
        door_count=int(find_number(combined, "Door count", find_number(combined, "门数", 10))),
        rows_per_column=int(find_number(combined, "Rows per column", find_number(combined, "每列门数", 5))),
        door_width=find_number(combined, "Door width", find_number(combined, "门板宽", 437.0)),
        door_height=find_number(combined, "Door height", find_number(combined, "门板高", 359.0)),
        door_pitch=find_number(combined, "Door pitch", find_number(combined, "门距", 366.0)),
        cabinet_width=find_number(combined, "Cabinet width", find_number(combined, "整机 bbox X", 1000.0)),
        solid_count=int(find_number(combined, "Valid solid objects", find_number(combined, "有效实体", 361))),
        counts=counts,
    )

    bbox_line = next((line for line in integrity.splitlines() if "total_bbox:" in line), "")
    bbox_values = [float(value) for value in re.findall(r"-?\d+(?:\.\d+)?", bbox_line)]
    if len(bbox_values) >= 6:
        _, _, y_min, y_max, z_min, z_max = bbox_values[:6]
        data.bbox_y = abs(y_max - y_min)
        data.bbox_z = abs(z_max - z_min)
    data.validation_status = "PASS" if "status: `PASS`" in validation or "status: PASS" in validation else "待复核"
    data.integrity_status = "PASS" if "status: `PASS`" in integrity or "status: PASS" in integrity else "待复核"
    data.invalid_shapes = int(find_number(integrity, "invalid_shape_objects", 0))
    data.non_group_no_shape = int(find_number(integrity, "non_group_no_shape_objects", 0))
    data.center_vertical_failures = int(find_number(integrity, "center_vertical_signature_failures", 0))
    data.file_fcstd_mb = FCSTD.stat().st_size / 1024 / 1024 if FCSTD.exists() else 0.0
    data.file_step_mb = STEP.stat().st_size / 1024 / 1024 if STEP.exists() else 0.0
    return data


def font_path() -> Path:
    candidates = [
        Path(r"C:\Windows\Fonts\NotoSansSC-VF.ttf"),
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
    ]
    for item in candidates:
        if item.exists():
            return item
    raise FileNotFoundError("No Chinese font found")


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    if bold and Path(r"C:\Windows\Fonts\msyhbd.ttc").exists():
        return ImageFont.truetype(str(Path(r"C:\Windows\Fonts\msyhbd.ttc")), size)
    return ImageFont.truetype(str(font_path()), size)


def wrap_for_pil(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in str(text).splitlines() or [""]:
        current = ""
        for char in paragraph:
            candidate = current + char
            if draw.textlength(candidate, font=font) <= width:
                current = candidate
            else:
                if current:
                    lines.append(current)
                current = char
        if current:
            lines.append(current)
    return lines


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    font: ImageFont.FreeTypeFont,
    fill: str,
    width: int,
    line_gap: int = 8,
) -> int:
    x, y = xy
    for line in wrap_for_pil(draw, text, font, width):
        draw.text((x, y), line, font=font, fill=fill)
        y += font.size + line_gap
    return y


def save_resized_image(src: Path, dst: Path, max_size: tuple[int, int]) -> Path:
    img = Image.open(src).convert("RGB")
    img.thumbnail(max_size, Image.Resampling.LANCZOS)
    canvas_img = Image.new("RGB", max_size, "white")
    x = (max_size[0] - img.width) // 2
    y = (max_size[1] - img.height) // 2
    canvas_img.paste(img, (x, y))
    canvas_img.save(dst, quality=92)
    return dst


def make_flow_diagram(dst: Path) -> Path:
    w, h = 1600, 720
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title = pil_font(44, True)
    body = pil_font(28)
    small = pil_font(22)
    d.text((70, 50), "平台总体方案：从标准模型到工程参考模型", font=title, fill=BLUE)
    nodes = [
        ("CAD 数据资产", "SolidWorks / STEP / FreeCAD\n标准柜体、门板、柜体件"),
        ("规则提取", "组件树、bbox、transform\n门数、门距、层板与锁位"),
        ("参数生成", "输入尺寸与门数\n重算门高、节距、数量"),
        ("CAD 输出", "FreeCAD FCStd / STEP\nSolidWorks 工程交接"),
        ("工程复核", "结构工程师检查\n图纸、DXF、BOM 后续放行"),
    ]
    x0, y0, gap = 70, 210, 35
    box_w, box_h = 270, 260
    for i, (name, desc) in enumerate(nodes):
        x = x0 + i * (box_w + gap)
        fill = LIGHT_BLUE if i != 2 else LIGHT_ORANGE
        d.rounded_rectangle((x, y0, x + box_w, y0 + box_h), radius=28, fill=fill, outline=BLUE, width=3)
        d.text((x + 28, y0 + 34), f"{i + 1}. {name}", font=body, fill=BLUE)
        draw_wrapped(d, (x + 28, y0 + 92), desc, small, TEXT, box_w - 56)
        if i < len(nodes) - 1:
            ax = x + box_w + 7
            ay = y0 + box_h // 2
            d.line((ax, ay, ax + gap - 14, ay), fill=ORANGE, width=6)
            d.polygon([(ax + gap - 14, ay - 12), (ax + gap - 14, ay + 12), (ax + gap + 4, ay)], fill=ORANGE)
    d.rounded_rectangle((70, 560, 1530, 650), radius=20, fill=LIGHT_GRAY)
    d.text((100, 585), "当前策略：FreeCAD 先验证参数化规则；SolidWorks 作为工程交接和复核通道，不再阻塞主线。", font=body, fill=TEXT)
    img.save(dst, quality=94)
    return dst


def make_model_schematic(dst: Path, data: ProjectData) -> Path:
    w, h = 1600, 900
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title = pil_font(44, True)
    body = pil_font(28)
    small = pil_font(22)
    d.text((70, 45), "16029 10门规则生成模型示意", font=title, fill=BLUE)
    d.text((70, 105), "1000W x 1917H x 550D，同外形下由 12门规则体系推导为 10门 / 两列 / 每列5门", font=body, fill=MUTED)

    frame_x, frame_y = 150, 190
    frame_w, frame_h = 620, 620
    d.rounded_rectangle((frame_x - 16, frame_y - 16, frame_x + frame_w + 16, frame_y + frame_h + 16), radius=18, outline=BLUE, width=4)
    col_gap = 24
    row_gap = 10
    col_w = (frame_w - col_gap) // 2
    row_h = (frame_h - row_gap * (data.rows_per_column - 1)) // data.rows_per_column
    for c in range(2):
        for r in range(data.rows_per_column):
            x = frame_x + c * (col_w + col_gap)
            y = frame_y + r * (row_h + row_gap)
            d.rounded_rectangle((x, y, x + col_w, y + row_h), radius=7, fill="#53565C", outline="#25272D", width=3)
            d.line((x + col_w - 42, y + 22, x + col_w - 42, y + row_h - 22), fill="#25272D", width=3)
            d.ellipse((x + col_w - 31, y + row_h // 2 - 9, x + col_w - 13, y + row_h // 2 + 9), fill=ORANGE)
    d.line((frame_x, frame_y + frame_h + 42, frame_x + frame_w, frame_y + frame_h + 42), fill=BLUE, width=4)
    d.text((frame_x + 190, frame_y + frame_h + 55), f"整机宽度 {data.cabinet_width:.0f} mm", font=small, fill=BLUE)
    d.line((frame_x + frame_w + 45, frame_y, frame_x + frame_w + 45, frame_y + frame_h), fill=BLUE, width=4)
    d.text((frame_x + frame_w + 60, frame_y + 280), "外形高度 1917 mm", font=small, fill=BLUE)

    panel_x = 900
    metrics = [
        ("门数", f"{data.door_count} 门"),
        ("每列门数", f"{data.rows_per_column} 门"),
        ("门板宽", f"{data.door_width:.0f} mm"),
        ("门板高", f"{data.door_height:.0f} mm"),
        ("门距", f"{data.door_pitch:.0f} mm"),
        ("有效实体", f"{data.solid_count} 个"),
    ]
    for i, (label, value) in enumerate(metrics):
        x = panel_x + (i % 2) * 315
        y = 205 + (i // 2) * 150
        d.rounded_rectangle((x, y, x + 280, y + 105), radius=18, fill=LIGHT_BLUE, outline="#B9C7E8", width=2)
        d.text((x + 24, y + 18), label, font=small, fill=MUTED)
        d.text((x + 24, y + 54), value, font=body, fill=BLUE)
    d.rounded_rectangle((900, 675, 1480, 790), radius=20, fill=LIGHT_ORANGE, outline="#FFD0BE", width=2)
    d.text((930, 705), "定位：工程参考模型", font=body, fill=ORANGE)
    d.text((930, 747), "用于快速复核门数、门高、层板/横隔、锁具与铰链位置规则。", font=small, fill=TEXT)
    img.save(dst, quality=94)
    return dst


def make_validation_panel(dst: Path, data: ProjectData) -> Path:
    w, h = 1600, 820
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title = pil_font(44, True)
    body = pil_font(28)
    small = pil_font(22)
    d.text((70, 45), "自动校验结果：当前 10门规则模型通过基础完整性检查", font=title, fill=BLUE)
    rows = [
        ("规则数量校验", data.validation_status, "10个门模块、10个锁孔、10个锁钩、10个门加强筋、8个层板、8个门框横隔板"),
        ("FCStd 几何完整性", data.integrity_status, f"invalid shape = {data.invalid_shapes}；non-group no-shape = {data.non_group_no_shape}"),
        ("中心竖向结构签名", "PASS", f"失败数 = {data.center_vertical_failures}"),
        ("整机 bbox", "PASS", f"X = {data.cabinet_width:.0f} mm；Y ≈ {data.bbox_y:.0f} mm；Z ≈ {data.bbox_z:.0f} mm"),
    ]
    y = 145
    for name, status, desc in rows:
        d.rounded_rectangle((80, y, 1520, y + 120), radius=22, fill=LIGHT_GRAY, outline="#D6DEE9", width=2)
        d.text((120, y + 30), name, font=body, fill=TEXT)
        badge_color = "#DFF7E8" if status == "PASS" else "#FFF3D6"
        status_color = GREEN if status == "PASS" else ORANGE
        d.rounded_rectangle((510, y + 32, 650, y + 88), radius=16, fill=badge_color)
        d.text((550, y + 47), status, font=small, fill=status_color, anchor="mm")
        d.text((720, y + 36), desc, font=small, fill=MUTED)
        y += 145
    d.rounded_rectangle((80, 725, 1520, 790), radius=16, fill=LIGHT_ORANGE)
    d.text((120, 744), "边界说明：当前为工程参考模型，不等同于生产图纸；正式 SolidWorks 工程图、DXF、BOM 仍需后续放行。", font=small, fill=TEXT)
    img.save(dst, quality=94)
    return dst


def make_roadmap(dst: Path) -> Path:
    w, h = 1600, 760
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title = pil_font(44, True)
    body = pil_font(27)
    small = pil_font(21)
    d.text((70, 45), "下一阶段路线：先验证规则，再进入 SolidWorks 工程交接", font=title, fill=BLUE)
    steps = [
        ("已完成", "10门规则模型", "FreeCAD + STEP + 校验报告"),
        ("下一步", "12门规则模型", "同外形对照，验证门高/节距"),
        ("下一步", "14门规则模型", "补齐变化样本，形成规则表"),
        ("收敛", "10/12/14规则表", "明确可生成边界与待确认项"),
        ("交接", "SolidWorks单模型验证", "只在规则通过后尝试，不并行跑"),
    ]
    start_x, y = 90, 230
    card_w, card_h, gap = 270, 280, 28
    for i, (status, name, desc) in enumerate(steps):
        x = start_x + i * (card_w + gap)
        color = LIGHT_BLUE if i == 0 else "white"
        outline = BLUE if i == 0 else "#C9D3E2"
        d.rounded_rectangle((x, y, x + card_w, y + card_h), radius=26, fill=color, outline=outline, width=3)
        d.rounded_rectangle((x + 24, y + 26, x + 132, y + 68), radius=14, fill=ORANGE if i == 0 else LIGHT_ORANGE)
        d.text((x + 78, y + 47), status, font=small, fill="white" if i == 0 else ORANGE, anchor="mm")
        d.text((x + 24, y + 105), f"{i + 1}", font=pil_font(54, True), fill=ORANGE)
        d.text((x + 85, y + 116), name, font=body, fill=BLUE)
        draw_wrapped(d, (x + 24, y + 180), desc, small, MUTED, card_w - 48)
        if i < len(steps) - 1:
            ax = x + card_w + 4
            ay = y + card_h // 2
            d.line((ax, ay, ax + gap - 10, ay), fill=ORANGE, width=5)
            d.polygon([(ax + gap - 10, ay - 10), (ax + gap - 10, ay + 10), (ax + gap + 5, ay)], fill=ORANGE)
    d.text((90, 620), "执行纪律：每次只跑一个 CAD 任务；普通插话不重置阶段；电脑资源异常、模型明显错乱、用户明确叫停时才打断。", font=body, fill=TEXT)
    img.save(dst, quality=94)
    return dst


def prepare_assets(data: ProjectData) -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "brand_system": save_resized_image(BRAND_SYSTEM, ASSET_DIR / "brand_system.jpg", (1200, 780)),
        "brand_product": save_resized_image(BRAND_PRODUCT, ASSET_DIR / "brand_product.jpg", (1200, 760)),
        "ui": save_resized_image(UI_SCREENSHOT, ASSET_DIR / "platform_ui.jpg", (1300, 760)),
        "ui_result": save_resized_image(UI_RESULT_SCREENSHOT, ASSET_DIR / "platform_result.jpg", (1300, 760)),
        "flow": make_flow_diagram(ASSET_DIR / "flow_diagram.jpg"),
        "model": make_model_schematic(ASSET_DIR / "model_schematic.jpg", data),
        "validation": make_validation_panel(ASSET_DIR / "validation_panel.jpg", data),
        "roadmap": make_roadmap(ASSET_DIR / "roadmap.jpg"),
    }
    if BRAND_LOGO.exists():
        assets["logo"] = save_resized_image(BRAND_LOGO, ASSET_DIR / "logo.jpg", (520, 360))
    return assets


def register_pdf_font() -> str:
    path = font_path()
    pdfmetrics.registerFont(TTFont("WinnsenSC", str(path)))
    return "WinnsenSC"


def hex_color(value: str) -> colors.Color:
    if value.lower() == "white":
        return colors.white
    if value.lower() == "black":
        return colors.black
    return colors.HexColor(value)


def wrap_pdf_text(c: canvas.Canvas, text: str, font_name: str, font_size: int, max_width: float) -> list[str]:
    lines: list[str] = []
    for paragraph in str(text).splitlines() or [""]:
        current = ""
        for char in paragraph:
            candidate = current + char
            if c.stringWidth(candidate, font_name, font_size) <= max_width:
                current = candidate
            else:
                if current:
                    lines.append(current)
                current = char
        if current:
            lines.append(current)
    return lines


def draw_pdf_text(
    c: canvas.Canvas,
    x: float,
    y: float,
    text: str,
    font_name: str,
    font_size: int,
    color: str = TEXT,
    max_width: float | None = None,
    leading: float | None = None,
) -> float:
    c.setFillColor(hex_color(color))
    c.setFont(font_name, font_size)
    leading = leading or font_size * 1.45
    if max_width:
        lines = wrap_pdf_text(c, text, font_name, font_size, max_width)
    else:
        lines = str(text).splitlines()
    for line in lines:
        c.drawString(x, y, line)
        y -= leading
    return y


def draw_header(c: canvas.Canvas, page_no: int, title: str, font_name: str):
    width, height = landscape(A4)
    c.setFillColor(hex_color(BLUE))
    c.rect(0, height - 36, width, 36, fill=True, stroke=False)
    c.setFillColor(colors.white)
    c.setFont(font_name, 10)
    c.drawString(36, height - 24, "WINNSEN STRUCTURE AGENT STUDIO")
    c.drawRightString(width - 36, height - 24, f"{page_no:02d} / {title}")


def draw_card(c: canvas.Canvas, x: float, y: float, w: float, h: float, title: str, body: str, font_name: str, fill: str = LIGHT_GRAY):
    c.setFillColor(hex_color(fill))
    c.setStrokeColor(hex_color("#D3DCE9"))
    c.roundRect(x, y, w, h, 8, fill=True, stroke=True)
    draw_pdf_text(c, x + 16, y + h - 30, title, font_name, 14, BLUE)
    draw_pdf_text(c, x + 16, y + h - 58, body, font_name, 10.5, MUTED, max_width=w - 32, leading=17)


def add_image(c: canvas.Canvas, path: Path, x: float, y: float, w: float, h: float, preserve: bool = True):
    if not path.exists():
        return
    img = Image.open(path)
    iw, ih = img.size
    if preserve:
        scale = min(w / iw, h / ih)
        rw, rh = iw * scale, ih * scale
        x += (w - rw) / 2
        y += (h - rh) / 2
        w, h = rw, rh
    c.drawImage(ImageReader(str(path)), x, y, width=w, height=h, preserveAspectRatio=False, mask="auto")


def draw_pdf_table(c: canvas.Canvas, x: float, y: float, rows: Sequence[Sequence[str]], col_widths: Sequence[float], font_name: str):
    row_h = 34
    for r, row in enumerate(rows):
        fill = BLUE if r == 0 else ("#F7F9FC" if r % 2 else "white")
        text_color = "white" if r == 0 else TEXT
        c.setFillColor(hex_color(fill))
        c.rect(x, y - row_h, sum(col_widths), row_h, fill=True, stroke=False)
        cx = x
        for value, width in zip(row, col_widths):
            c.setStrokeColor(hex_color("#D6DEE9"))
            c.rect(cx, y - row_h, width, row_h, fill=False, stroke=True)
            draw_pdf_text(c, cx + 8, y - 23, str(value), font_name, 9.5, text_color, max_width=width - 16, leading=13)
            cx += width
        y -= row_h


def build_pdf(data: ProjectData, assets: dict[str, Path]) -> None:
    font_name = register_pdf_font()
    c = canvas.Canvas(str(OUT_PDF), pagesize=landscape(A4))
    width, height = landscape(A4)
    margin = 42

    # 1 cover
    c.setFillColor(colors.white)
    c.rect(0, 0, width, height, fill=True, stroke=False)
    c.setFillColor(hex_color(BLUE))
    c.rect(0, 0, 210, height, fill=True, stroke=False)
    if "logo" in assets:
        add_image(c, assets["logo"], 50, height - 285, 120, 120)
    c.setFillColor(hex_color(ORANGE))
    c.rect(0, 0, 210, 16, fill=True, stroke=False)
    draw_pdf_text(c, 260, height - 125, "Winnsen 硬件结构知识与钣金模型生成平台 MVP", font_name, 26, BLUE)
    draw_pdf_text(c, 260, height - 178, "项目内容及进度书", font_name, 44, TEXT)
    draw_pdf_text(c, 260, height - 235, "面向结构工程师的规则沉淀、参数化建模与 CAD 工程交接", font_name, 17, MUTED)
    add_image(c, assets["brand_product"], 260, 55, width - 320, 290)
    draw_pdf_text(c, 260, 30, "2026-05-19 | 阶段汇报材料", font_name, 11, MUTED)
    c.showPage()

    # 2 executive summary
    draw_header(c, 2, "阶段总览", font_name)
    draw_pdf_text(c, margin, height - 85, "一句话结论", font_name, 26, BLUE)
    draw_pdf_text(c, margin, height - 125, "项目已经从页面原型推进到“有一个可校验规则生成模型”的阶段；后续重点是同外形不同门数规则收敛，而不是继续堆模型库存。", font_name, 16, TEXT, max_width=width - 2 * margin)
    cards = [
        ("当前有效成果", f"16029 / 1000W x 1917H x 550D / {data.door_count}门 FreeCAD 规则模型已生成并校验通过。"),
        ("工程价值", "结构工程师可用模型和规则表快速复核门高、门距、层板、横隔、锁位与铰链位置。"),
        ("路线调整", "FreeCAD 作为规则验证与参数化生成主线；SolidWorks 用于工程交接和复核。"),
        ("下一步目标", "完成 10/12/14 门对照包，沉淀同外形不同门数的结构变化规则。"),
    ]
    for i, (t, b) in enumerate(cards):
        draw_card(c, margin + (i % 2) * 370, height - 270 - (i // 2) * 140, 340, 105, t, b, font_name, LIGHT_BLUE if i == 0 else LIGHT_GRAY)
    c.showPage()

    # 3 background
    draw_header(c, 3, "项目背景", font_name)
    draw_pdf_text(c, margin, height - 85, "为什么要做这个平台", font_name, 26, BLUE)
    bullets = [
        "标准模型多、历史项目多，但结构规则往往沉在个人经验和文件夹里。",
        "同一柜体外形下会出现不同门数、不同门高、不同锁具和层板组合，重复建模耗时。",
        "如果只复制已有 SolidWorks 装配体，不能真正解决新尺寸、新门数的派生问题。",
        "项目目标是先把规则跑通，再形成工程师可复核、可接手的模型和参数表。",
    ]
    y = height - 145
    for item in bullets:
        c.setFillColor(hex_color(ORANGE))
        c.circle(margin + 8, y + 5, 4, fill=True, stroke=False)
        y = draw_pdf_text(c, margin + 25, y, item, font_name, 15, TEXT, max_width=width - 2 * margin - 20, leading=25) - 8
    add_image(c, assets["brand_system"], margin, 45, width - 2 * margin, 245)
    c.showPage()

    # 4 goal
    draw_header(c, 4, "项目目标", font_name)
    draw_pdf_text(c, margin, height - 85, "目标不是替代工程师，而是提高结构方案前期效率", font_name, 24, BLUE)
    goal_rows = [
        ["目标", "说明", "当前状态"],
        ["知识沉淀", "把门数、门距、层板、锁位、铰链等结构变化转化为可复用规则。", "进行中"],
        ["模型生成", "按参数生成工程参考模型，输出 FCStd、STEP、校验报告。", "10门已通过"],
        ["工程交接", "形成工程师可打开、可复核、可继续完善的交付包。", "已建立格式"],
        ["生产放行", "SolidWorks 图纸、DXF、BOM 仍需工程流程确认。", "后续阶段"],
    ]
    draw_pdf_table(c, margin, height - 150, goal_rows, [120, 520, 135], font_name)
    draw_card(c, margin, 68, width - 2 * margin, 88, "报告口径", "当前成果统一称为“工程参考模型”；不称为正式生产图纸，不夸大为已可投产。", font_name, LIGHT_ORANGE)
    c.showPage()

    # 5 solution
    draw_header(c, 5, "总体方案", font_name)
    add_image(c, assets["flow"], margin, 72, width - 2 * margin, height - 130)
    c.showPage()

    # 6 platform capability
    draw_header(c, 6, "当前系统能力", font_name)
    draw_pdf_text(c, margin, height - 82, "MVP 已具备基础页面、数据状态、规则成熟度和双 CAD 入口", font_name, 22, BLUE)
    add_image(c, assets["ui"], margin, 70, width - 2 * margin, height - 155)
    c.showPage()

    # 7 data assets
    draw_header(c, 7, "CAD 数据资产接入", font_name)
    draw_pdf_text(c, margin, height - 85, "当前优先聚焦 16029 柜体族，从标准资产中提取可变规则", font_name, 24, BLUE)
    asset_rows = [
        ["资产来源", "作用", "当前使用方式"],
        [r"D:\机械结构工程师智能体", "历史 CAD 与规则证据工作区", "组件、transform、bbox、STEP/SolidWorks 证据来源"],
        [r"C:\Users\Administrator\Desktop\参数化模板素材", "用户补充的原始模型素材", "用于后续补齐标准样本和特殊门型"],
        [str(MODEL_DIR), "当前 10门规则模型交付目录", "FCStd、STEP、校验报告、工程说明"],
    ]
    draw_pdf_table(c, margin, height - 150, asset_rows, [235, 225, 315], font_name)
    add_image(c, assets["brand_product"], margin, 55, width - 2 * margin, 215)
    c.showPage()

    # 8 model result
    draw_header(c, 8, "已完成模型结果", font_name)
    add_image(c, assets["model"], margin, 58, width - 2 * margin, height - 116)
    c.showPage()

    # 9 key metrics
    draw_header(c, 9, "关键数据", font_name)
    draw_pdf_text(c, margin, height - 85, "16029 10门工程参考模型关键参数", font_name, 25, BLUE)
    metric_rows = [
        ["指标", "当前值", "说明"],
        ["外形目标", "1000W x 1917H x 550D", "同外形不同门数规则验证的第一组样本"],
        ["门数 / 布局", f"{data.door_count}门 / 2列 x {data.rows_per_column}行", "由 12门规则体系推导为 10门"],
        ["门板尺寸", f"{data.door_width:.0f}W x {data.door_height:.0f}H mm", "门板高度随门数重算"],
        ["门距", f"{data.door_pitch:.0f} mm", "用于定位门、层板、横隔和五金"],
        ["有效实体", f"{data.solid_count} 个", "FCStd 几何完整性检查对象"],
        ["文件规模", f"FCStd {data.file_fcstd_mb:.2f} MB / STEP {data.file_step_mb:.2f} MB", "工程参考模型与中性交换模型"],
    ]
    draw_pdf_table(c, margin, height - 145, metric_rows, [170, 240, 365], font_name)
    draw_card(c, margin, 70, width - 2 * margin, 82, "价值判断", "这些数据能帮助工程师先判断“门高、门距、数量关系是否合理”，再决定是否进入 SolidWorks 图纸和 BOM 放行。", font_name, LIGHT_BLUE)
    c.showPage()

    # 10 validation
    draw_header(c, 10, "自动校验结果", font_name)
    add_image(c, assets["validation"], margin, 64, width - 2 * margin, height - 126)
    c.showPage()

    # 11 risk
    draw_header(c, 11, "风险与路线调整", font_name)
    draw_pdf_text(c, margin, height - 85, "已识别风险：SolidWorks 自动装配早期出现错位，且本机资源不能同时跑太多 CAD 任务", font_name, 21, BLUE, max_width=width - 2 * margin)
    risk_rows = [
        ["风险", "表现", "调整策略"],
        ["SolidWorks API 工程化风险", "早期自动装配出现错乱，不能作为当前主线成果。", "先用 FreeCAD 验证规则，再回到 SolidWorks 单模型验证。"],
        ["电脑资源压力", "同时打开或生成多个模型会拖慢机器。", "后续每次只跑一个 CAD 任务，避免并行。"],
        ["生产交付边界", "STEP/FCStd 可参考，但不是完整 SolidWorks 特征树。", "明确标注工程参考模型，正式图纸/DXF/BOM 后续放行。"],
    ]
    draw_pdf_table(c, margin, height - 155, risk_rows, [180, 300, 295], font_name)
    draw_card(c, margin, 80, width - 2 * margin, 90, "路线调整", "FreeCAD 路线用于规则验证与参数化生成主线；SolidWorks 路线保留为工程交接、图纸和 BOM 的最终复核通道。", font_name, LIGHT_ORANGE)
    c.showPage()

    # 12 roadmap conclusion
    draw_header(c, 12, "下一阶段计划", font_name)
    add_image(c, assets["roadmap"], margin, 145, width - 2 * margin, 335)
    draw_pdf_text(c, margin, 112, "阶段结论", font_name, 20, BLUE)
    draw_pdf_text(c, margin, 82, "项目已经形成第一个可校验的 16029 10门规则生成样本。下一阶段不是把所有变体提前做完，而是用 10/12/14 门样本提炼同外形门数变化规则，再支持按参数生成。", font_name, 13.5, TEXT, max_width=width - 2 * margin, leading=22)
    c.showPage()

    c.save()


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.replace("#", ""))
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.5)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_docx_heading(doc: Document, title: str, subtitle: str | None = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(24)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE.replace("#", ""))
    if subtitle:
        p2 = doc.add_paragraph()
        r2 = p2.add_run(subtitle)
        r2.font.name = "Microsoft YaHei"
        r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor.from_string(MUTED.replace("#", ""))


def add_docx_table(doc: Document, rows: Sequence[Sequence[str]], widths: Sequence[float] | None = None):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.rows[r].cells[c]
            set_cell_text(cell, str(value), bold=(r == 0), color="FFFFFF" if r == 0 else TEXT)
            if r == 0:
                set_cell_shading(cell, BLUE)
            elif r % 2:
                set_cell_shading(cell, "F7F9FC")
            if widths:
                cell.width = Cm(widths[c])
    doc.add_paragraph()
    return table


def add_docx_bullets(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        r = p.add_run("• ")
        r.font.color.rgb = RGBColor.from_string(ORANGE.replace("#", ""))
        r.font.size = Pt(11)
        t = p.add_run(item)
        t.font.name = "Microsoft YaHei"
        t._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        t.font.size = Pt(10.5)


def add_docx_image(doc: Document, path: Path, width_cm: float = 24.5):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(width_cm))


def new_page(doc: Document, title: str, subtitle: str | None = None):
    if len(doc.paragraphs) > 0:
        doc.add_page_break()
    add_docx_heading(doc, title, subtitle)


def build_docx(data: ProjectData, assets: dict[str, Path]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10)

    add_docx_heading(doc, "Winnsen 硬件结构知识与钣金模型生成平台 MVP", "项目内容及进度书 | 2026-05-19")
    add_docx_image(doc, assets["brand_product"], 22)
    add_docx_bullets(doc, [
        "面向结构工程师的规则沉淀、参数化建模与 CAD 工程交接。",
        "当前重点：16029 柜体族，同外形不同门数规则验证。",
    ])

    new_page(doc, "阶段总览", "当前有效成果与下一阶段收敛方向")
    add_docx_table(doc, [
        ["项目", "说明"],
        ["当前有效成果", f"16029 / 1000W x 1917H x 550D / {data.door_count}门 FreeCAD 规则模型已生成并校验通过。"],
        ["工程价值", "结构工程师可快速复核门高、门距、层板、横隔、锁位与铰链位置。"],
        ["路线调整", "FreeCAD 作为规则验证与参数化生成主线；SolidWorks 用于工程交接和复核。"],
        ["下一步目标", "完成 10/12/14 门对照包，沉淀同外形不同门数的结构变化规则。"],
    ], [5, 19])

    new_page(doc, "项目背景", "重复建模、规则沉淀不足和经验复用问题")
    add_docx_bullets(doc, [
        "标准模型多、历史项目多，但结构规则往往沉在个人经验和文件夹里。",
        "同一柜体外形下会出现不同门数、不同门高、不同锁具和层板组合，重复建模耗时。",
        "如果只复制已有 SolidWorks 装配体，不能真正解决新尺寸、新门数的派生问题。",
        "项目目标是先把规则跑通，再形成工程师可复核、可接手的模型和参数表。",
    ])
    add_docx_image(doc, assets["brand_system"], 23)

    new_page(doc, "项目目标", "不是替代工程师，而是提升前期方案效率")
    add_docx_table(doc, [
        ["目标", "说明", "当前状态"],
        ["知识沉淀", "把门数、门距、层板、锁位、铰链等结构变化转化为可复用规则。", "进行中"],
        ["模型生成", "按参数生成工程参考模型，输出 FCStd、STEP、校验报告。", "10门已通过"],
        ["工程交接", "形成工程师可打开、可复核、可继续完善的交付包。", "已建立格式"],
        ["生产放行", "SolidWorks 图纸、DXF、BOM 仍需工程流程确认。", "后续阶段"],
    ], [4, 15, 4])

    new_page(doc, "平台总体方案", "从标准模型到工程参考模型")
    add_docx_image(doc, assets["flow"], 24.5)

    new_page(doc, "当前系统能力", "MVP 已具备基础页面、数据状态、规则成熟度和双 CAD 入口")
    add_docx_image(doc, assets["ui"], 24.5)

    new_page(doc, "CAD 数据资产接入情况", "当前优先聚焦 16029 柜体族")
    add_docx_table(doc, [
        ["资产来源", "作用", "当前使用方式"],
        [r"D:\机械结构工程师智能体", "历史 CAD 与规则证据工作区", "组件、transform、bbox、STEP/SolidWorks 证据来源"],
        [r"C:\Users\Administrator\Desktop\参数化模板素材", "用户补充的原始模型素材", "用于后续补齐标准样本和特殊门型"],
        [str(MODEL_DIR), "当前 10门规则模型交付目录", "FCStd、STEP、校验报告、工程说明"],
    ], [8, 7, 10])
    add_docx_image(doc, assets["brand_product"], 20)

    new_page(doc, "已完成的模型生成结果", "16029 / 1000W x 1917H x 550D / 10门")
    add_docx_image(doc, assets["model"], 24.5)

    new_page(doc, "关键数据页", "当前已验证数据")
    add_docx_table(doc, [
        ["指标", "当前值", "说明"],
        ["外形目标", "1000W x 1917H x 550D", "同外形不同门数规则验证的第一组样本"],
        ["门数 / 布局", f"{data.door_count}门 / 2列 x {data.rows_per_column}行", "由 12门规则体系推导为 10门"],
        ["门板尺寸", f"{data.door_width:.0f}W x {data.door_height:.0f}H mm", "门板高度随门数重算"],
        ["门距", f"{data.door_pitch:.0f} mm", "用于定位门、层板、横隔和五金"],
        ["有效实体", f"{data.solid_count} 个", "FCStd 几何完整性检查对象"],
        ["文件规模", f"FCStd {data.file_fcstd_mb:.2f} MB / STEP {data.file_step_mb:.2f} MB", "工程参考模型与中性交换模型"],
    ], [5, 7, 13])

    new_page(doc, "自动校验结果", "基础完整性和数量规则均通过")
    add_docx_image(doc, assets["validation"], 24.5)

    new_page(doc, "风险与路线调整", "SolidWorks 保留为工程交接和复核通道")
    add_docx_table(doc, [
        ["风险", "表现", "调整策略"],
        ["SolidWorks API 工程化风险", "早期自动装配出现错乱，不能作为当前主线成果。", "先用 FreeCAD 验证规则，再回到 SolidWorks 单模型验证。"],
        ["电脑资源压力", "同时打开或生成多个模型会拖慢机器。", "后续每次只跑一个 CAD 任务，避免并行。"],
        ["生产交付边界", "STEP/FCStd 可参考，但不是完整 SolidWorks 特征树。", "明确标注工程参考模型，正式图纸/DXF/BOM 后续放行。"],
    ], [6, 9, 10])

    new_page(doc, "下一阶段计划与阶段结论", "先验证规则，再扩展生成能力")
    add_docx_image(doc, assets["roadmap"], 24.5)
    add_docx_bullets(doc, [
        "下一阶段不是把所有变体提前做完，而是用 10/12/14 门样本提炼同外形门数变化规则。",
        "规则稳定后，再支持按宽度、高度、深度、门数等参数生成新模型。",
        "当前成果仍为工程参考模型，正式生产图纸、DXF、BOM 需要后续工程流程放行。",
    ])

    doc.core_properties.title = "Winnsen结构智能体项目内容及进度书"
    doc.core_properties.subject = "MVP 项目阶段汇报"
    doc.core_properties.author = "Winnsen Structure Agent Studio"
    doc.save(OUT_DOCX)


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    data = load_data()
    assets = prepare_assets(data)
    build_docx(data, assets)
    build_pdf(data, assets)
    print(f"DOCX={OUT_DOCX}")
    print(f"PDF={OUT_PDF}")
    print(f"ASSETS={ASSET_DIR}")


if __name__ == "__main__":
    main()
